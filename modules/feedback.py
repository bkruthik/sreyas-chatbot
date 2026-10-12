"""
modules/feedback.py
================================================================================
Feedback Storage Module.
Persists user ratings (👍 Yes / 👎 No) into a dedicated `feedback.xlsx` workbook,
completely separate from lead enquiries.
================================================================================
"""

import os
import threading
import logging
from datetime import datetime
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import config

logger = logging.getLogger("sreyas_chatbot.feedback")

SHEET_NAME = "Feedback"
COLUMNS = ["Feedback ID", "Date", "Time", "Rating", "Category", "Session ID"]
_ID_PATTERN_PREFIX = "FB"


class FeedbackManager:
    """Manages feedback collection and analytics in feedback.xlsx."""

    def __init__(self, excel_path: str = config.FEEDBACK_FILE) -> None:
        self.excel_path = excel_path
        self._lock = threading.Lock()
        self._ensure_workbook_exists()

    def _ensure_workbook_exists(self) -> None:
        """Create feedback.xlsx with header row if missing."""
        try:
            directory = os.path.dirname(self.excel_path)
            if directory and not os.path.exists(directory):
                os.makedirs(directory, exist_ok=True)

            if not os.path.exists(self.excel_path):
                wb = Workbook()
                ws = wb.active
                ws.title = SHEET_NAME
                ws.append(COLUMNS)
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                col_widths = [14, 12, 10, 16, 22, 36]
                for idx, width in enumerate(col_widths, start=1):
                    ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
                wb.save(self.excel_path)
                wb.close()
                logger.info("Created feedback storage file at '%s'.", self.excel_path)
        except Exception as e:
            logger.error("Error creating feedback workbook: %s", e)

    def record_feedback(self, rating: str, category: str = "GENERAL", session_id: str = "") -> Dict[str, Any]:
        """
        Save user feedback rating (helpful / not_helpful or yes / no / positive / negative).

        Args:
            rating: "helpful" / "yes" / "positive" or "not_helpful" / "no" / "negative".
            category: Chat category at time of rating.
            session_id: Session identifier.
        """
        norm_rating = "Helpful" if rating.lower() in ("helpful", "yes", "true", "1", "positive") else "Unhelpful"
        now = datetime.now()

        with self._lock:
            try:
                self._ensure_workbook_exists()
                wb = load_workbook(self.excel_path)
                ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active

                # Generate ID
                row_count = ws.max_row
                feedback_id = f"FB{row_count:04d}"

                row = [
                    feedback_id,
                    now.strftime("%Y-%m-%d"),
                    now.strftime("%H:%M:%S"),
                    norm_rating,
                    category.upper(),
                    session_id,
                ]
                ws.append(row)
                wb.save(self.excel_path)
                wb.close()
                logger.info("Saved feedback '%s' (%s).", feedback_id, norm_rating)
                return {"success": True, "message": "Thank you for your feedback!"}
            except Exception as e:
                logger.error("Failed to save feedback: %s", e)
                return {"success": False, "message": "Could not record feedback."}

    def get_analytics_summary(self) -> Dict[str, Any]:
        """Compute aggregated feedback statistics for the admin dashboard."""
        with self._lock:
            if not os.path.exists(self.excel_path):
                return {
                    "total": 0, "helpful": 0, "unhelpful": 0,
                    "positive": 0, "negative": 0, "helpful_ratio": 100.0, "today_total": 0
                }

            try:
                wb = load_workbook(self.excel_path, data_only=True)
                ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
                records = list(ws.iter_rows(min_row=2, values_only=True))
                wb.close()

                total = 0
                helpful = 0
                unhelpful = 0
                today_str = datetime.now().strftime("%Y-%m-%d")
                today_total = 0

                for row in records:
                    if not row or not row[0]:
                        continue
                    total += 1
                    rating_val = str(row[3] or "").strip().title()
                    date_val = str(row[1] or "").strip()

                    if rating_val == "Helpful":
                        helpful += 1
                    elif rating_val == "Unhelpful":
                        unhelpful += 1

                    if date_val == today_str:
                        today_total += 1

                ratio = round((helpful / total * 100), 1) if total > 0 else 100.0

                return {
                    "total": total,
                    "helpful": helpful,
                    "unhelpful": unhelpful,
                    "positive": helpful,
                    "negative": unhelpful,
                    "helpful_ratio": ratio,
                    "today_total": today_total,
                }
            except Exception as e:
                logger.error("Error reading feedback summary: %s", e)
                return {
                    "total": 0, "helpful": 0, "unhelpful": 0,
                    "positive": 0, "negative": 0, "helpful_ratio": 100.0, "today_total": 0
                }
